"""Download into memory; parse supported files in disposable subprocesses."""
import hashlib
import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path


def xlsx_xml_fallback(data):
    """Salvage valid cell XML when Office-specific rich styles are rejected."""
    import posixpath
    from lxml import etree
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    lines, count = [], 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            root = etree.fromstring(archive.read('xl/sharedStrings.xml'), parser)
            strings = [''.join(si.xpath('.//s:t/text()', namespaces=ns)) for si in root]
        workbook = etree.fromstring(archive.read('xl/workbook.xml'), parser)
        relations = etree.fromstring(archive.read('xl/_rels/workbook.xml.rels'), parser)
        paths = {r.get('Id'): r.get('Target') for r in relations}
        for sheet in workbook.xpath('//s:sheet', namespaces=ns):
            rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = paths[rid]
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            lines.append('[工作表：' + sheet.get('name', '') + ']')
            cells = etree.fromstring(archive.read(path), parser)
            for row in cells.xpath('//s:sheetData/s:row', namespaces=ns):
                values = []
                for cell in row:
                    count += 1
                    if count > 2000000:
                        return {'status': 'partial', 'text': '\n'.join(lines), 'error': 'cell_limit_2000000'}
                    value = ''.join(cell.xpath('./s:v/text()', namespaces=ns))
                    formula = cell.xpath('./s:f/text()', namespaces=ns)
                    if formula:
                        value = '=' + formula[0]
                    elif cell.get('t') == 's':
                        value = strings[int(value)] if value else ''
                    elif cell.get('t') == 'inlineStr':
                        value = ''.join(cell.xpath('.//s:t/text()', namespaces=ns))
                    values.append(cell.get('r', '') + ': ' + value)
                lines.append('\t'.join(values))
    return {'status': 'partial', 'text': '\n'.join(lines), 'parser': 'xlsx_raw_xml',
            'error': 'nonstandard_workbook_styles_raw_cell_values_extracted'}


def parse_bytes(data, extension):
    if data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        return {'status': 'needs_manual_review', 'text': '', 'error': 'legacy_or_encrypted_ole_container',
                'detected_format': 'ole_binary'}
    if extension in ('docx', 'xlsx'):
        if not zipfile.is_zipfile(io.BytesIO(data)):
            return {'status': 'needs_manual_review', 'text': '', 'error': 'extension_mismatch_or_invalid_office_container'}
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 300 * 1024 * 1024:
                raise ValueError('Expanded Office file exceeds 300 MiB limit')
    if extension == 'pdf':
        import pypdfium2 as pdfium
        import pypdfium2.raw as raw
        try:
            reader = pdfium.PdfDocument(data)
        except Exception:
            from pypdf import PdfReader
            fallback = PdfReader(io.BytesIO(data))
            if fallback.is_encrypted and not fallback.decrypt(''):
                return {'status': 'needs_manual_review', 'text': '', 'error': 'encrypted_pdf'}
            raise
        texts, scanned = [], 0
        count = len(reader)
        try:
            for number in range(count):
                if number >= 600:
                    return {'status': 'partial', 'text': '\n\n'.join(texts), 'error': 'page_limit_600'}
                page = reader[number]
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded(errors='strict').strip()
                    if len(text) < 20 and next(page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE]), None) is not None:
                        scanned += 1
                    texts.append(f'[第 {number + 1} 页]\n{text}')
                finally:
                    textpage.close()
                    page.close()
        finally:
            reader.close()
        if scanned:
            has_text = any(len(t.split('\n', 1)[-1]) >= 20 for t in texts)
            return {'status': 'partial' if has_text else 'needs_manual_review',
                    'text': '\n\n'.join(texts) if has_text else '',
                    'error': f'image_pages_without_ocr:{scanned}', 'page_count': count, 'parser': 'pdfium2'}
        combined = '\n\n'.join(texts)
        if not any(t.split('\n', 1)[-1].strip() for t in texts):
            return {'status': 'needs_manual_review', 'text': '', 'error': 'no_extractable_text'}
        return {'status': 'parsed', 'text': combined, 'error': None, 'page_count': count, 'parser': 'pdfium2'}
    if extension == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        # Preserve tables in document order instead of appending every table last.
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        lines = []
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                lines.append(Paragraph(child, doc).text)
            elif child.tag.endswith('}tbl'):
                lines.extend('\t'.join(c.text for c in row.cells) for row in Table(child, doc).rows)
        text = '\n'.join(lines).strip()
        drawings = bool(doc.element.xpath('.//w:drawing'))
        return {'status': ('partial' if text else 'needs_manual_review') if drawings else ('parsed' if text else 'needs_manual_review'),
                'text': text, 'error': 'embedded_images_without_ocr' if drawings else (None if text else 'no_extractable_text')}
    if extension == 'xlsx':
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except ValueError:
            return xlsx_xml_fallback(data)
        lines = []
        count = 0
        try:
            for sheet in book:
                lines.append(f'[工作表：{sheet.title}]')
                for row in sheet.iter_rows(values_only=True):
                    count += len(row)
                    if count > 2000000:
                        return {'status': 'partial', 'text': '\n'.join(lines), 'error': 'cell_limit_2000000'}
                    text = '\t'.join('' if v is None else str(v) for v in row).rstrip()
                    if text:
                        lines.append(text)
        finally:
            book.close()
        return {'status': 'parsed', 'text': '\n'.join(lines), 'error': None}
    return {'status': 'needs_manual_review', 'text': '', 'error': 'unsupported_format_no_ocr'}


def process(client, attachment):
    result = dict(attachment)
    try:
        response = client.get(attachment['url'])
        data = response['body']
        result.update(size_bytes=len(data), sha256=hashlib.sha256(data).hexdigest(),
                      final_url=response['url'], content_type=response['headers'].get('content-type'))
        extension = attachment['format']
        if data.lstrip()[:30].lower().startswith((b'<!doctype html', b'<html')):
            result.update(status='needs_manual_review', format='html', declared_format=extension,
                          error='attachment_link_returns_html')
        elif extension not in ('pdf', 'docx', 'xlsx'):
            result.update(status='needs_manual_review', error='unsupported_format_no_ocr')
        else:
            with tempfile.TemporaryDirectory(prefix='policy-attachment-') as temp:
                path = Path(temp) / ('file.' + extension)
                path.write_bytes(data)
                completed = subprocess.run([sys.executable, '-m', 'collector.attachments', str(path), extension],
                                           capture_output=True, timeout=90,
                                           env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
                if completed.returncode:
                    raise ValueError(completed.stderr.decode('utf-8', 'replace')[-1200:])
                result.update(json.loads(completed.stdout))
    except subprocess.TimeoutExpired:
        result.update(status='failed', error='parse_timeout_90_seconds')
    except Exception as exc:
        result.update(status='failed', error=f'{type(exc).__name__}: {exc}')
    result['needs_manual_review'] = result['status'] != 'parsed'
    return result


if __name__ == '__main__':
    try:
        result = parse_bytes(Path(sys.argv[1]).read_bytes(), sys.argv[2])
    except Exception as exc:
        result = {'status': 'failed', 'text': '', 'error': f'{type(exc).__name__}: {exc}'}
    print(json.dumps(result, ensure_ascii=False))
